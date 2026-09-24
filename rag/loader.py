"""
多格式文档加载器
================
支持 PDF / Word(docx) / Excel(xlsx) / CSV / TXT / Markdown，
统一转成 Document 列表，metadata 里带文件名等溯源信息。
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path


@dataclass
class Document:
    content: str
    metadata: dict = field(default_factory=dict)


def _load_txt(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(content=text, metadata={"source": path.name})]


def _load_pdf(path: Path) -> list[Document]:
    from pypdf import PdfReader

    docs: list[Document] = []
    reader = PdfReader(str(path))
    for i, page in enumerate(reader.pages):
        text = page.extract_text() or ""
        if text.strip():
            docs.append(Document(
                content=text,
                metadata={"source": path.name, "page": i + 1},
            ))
    return docs


def _load_docx(path: Path) -> list[Document]:
    import docx

    d = docx.Document(str(path))
    paragraphs = [p.text for p in d.paragraphs if p.text.strip()]
    text = "\n".join(paragraphs)
    return [Document(content=text, metadata={"source": path.name})]


def _load_xlsx(path: Path) -> list[Document]:
    import openpyxl

    docs: list[Document] = []
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    for ws in wb.worksheets:
        rows = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None]
            if cells:
                rows.append(" | ".join(cells))
        if rows:
            docs.append(Document(
                content=f"[Sheet: {ws.title}]\n" + "\n".join(rows),
                metadata={"source": path.name, "sheet": ws.title},
            ))
    wb.close()
    return docs


def _load_csv(path: Path) -> list[Document]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [Document(content=text, metadata={"source": path.name})]


_LOADERS = {
    ".txt": _load_txt,
    ".md": _load_txt,
    ".markdown": _load_txt,
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".xlsx": _load_xlsx,
    ".csv": _load_csv,
}


def load_file(path: str | Path) -> list[Document]:
    """加载单个文件，返回 Document 列表。不支持的扩展名会报错。"""
    path = Path(path)
    ext = path.suffix.lower()
    loader = _LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"暂不支持的文件类型: {ext}（支持 {list(_LOADERS)}）")
    docs = loader(path)
    if not docs:
        raise ValueError(f"文件 {path.name} 未解析出任何文本内容")
    return docs


def load_dir(dir_path: str | Path) -> list[Document]:
    """加载目录下所有支持的文档。"""
    dir_path = Path(dir_path)
    all_docs: list[Document] = []
    for ext in _LOADERS:
        for p in dir_path.glob(f"*{ext}"):
            try:
                all_docs.extend(load_file(p))
            except Exception as e:  # noqa: BLE001
                print(f"[警告] 跳过 {p.name}: {e}")
    if not all_docs:
        raise ValueError(f"目录 {dir_path} 下没有可加载的文档")
    return all_docs
